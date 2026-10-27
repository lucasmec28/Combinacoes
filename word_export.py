"""Lista editável de combinações para memória de cálculo, sem equação normativa."""
from io import BytesIO
from decimal import Decimal
import re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from engine import FAMILIES, InputError
from project_io import combination_type


def word_text(value):
    """Remove apenas caracteres impossíveis em XML1.0; preserva texto Unicode."""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(value))


def coefficient_text(value):
    """No mínimo duas casas, sem arredondar precisão adicional do motor."""
    value = Decimal(str(value))
    if not value.is_finite():
        raise InputError('Coeficiente não finito na exportação Word.')
    text = format(abs(value), 'f')
    if '.' not in text:
        text += '.00'
    else:
        integer, fraction = text.split('.')
        text = integer + '.' + fraction.rstrip('0').ljust(2, '0')
    return text.replace('.', ',')


def combination_expression(combination):
    names = {int(d['case']): str(d['name']) for d in combination.detail}
    terms = []
    for case, raw in combination.cases:
        if case not in names:
            raise InputError(f'Nome do caso {case} ausente na combinação {combination.name}.')
        coefficient = Decimal(str(raw))
        if coefficient == 0:
            continue
        name = ' '.join(word_text(names[case]).split())
        if not name:
            raise InputError(f'Nome do caso {case} vazio.')
        term = coefficient_text(coefficient) + name
        if not terms:
            terms.append(('-' if coefficient < 0 else '') + term)
        else:
            terms.append((' - ' if coefficient < 0 else ' + ') + term)
    return ''.join(terms) or '0'


def combinations_docx(combinations, project_name=''):
    """Mantém ordem/coeficientes selecionados; renumera a lista MC de 1 em diante."""
    rows = list(combinations)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05
    for style_name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[style_name].font.name = 'Arial'
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Title'].font.size = Pt(16)
    doc.styles['Heading 1'].font.size = Pt(12)
    # Alguns templates locais incluem borda decorativa no estilo Title.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.add_paragraph('Combinações de ações', 'Title')
    if project_name:
        doc.add_paragraph(word_text(project_name))
    previous_family = None
    for index, combination in enumerate(rows, 1):
        combination_type(combination.family)  # valida família conhecida
        if combination.family != previous_family:
            doc.add_paragraph(FAMILIES[combination.family], 'Heading 1')
            previous_family = combination.family
        paragraph = doc.add_paragraph(f'{index}- {combination_expression(combination)}')
        paragraph.paragraph_format.left_indent = Cm(0.9)
        paragraph.paragraph_format.first_line_indent = Cm(-0.9)
        paragraph.paragraph_format.keep_together = True
    if not rows:
        doc.add_paragraph('Nenhuma combinação selecionada.')
    doc.core_properties.title = 'Combinações de ações'
    doc.core_properties.author = 'LRO Engenharia'
    doc.core_properties.subject = 'Lista de combinações selecionadas para memória de cálculo'
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
