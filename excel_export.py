"""Exportação do aplicativo para Excel com casos/coeficientes numéricos.
A aba ROBOT começa em A1, sem cabeçalho, para colar diretamente na coluna Nome.
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName


def robot_xlsx(combinations):
    wb=Workbook()
    ws=wb.active
    ws.title='ROBOT'
    ws.sheet_view.showGridLines=True
    ws.column_dimensions['A'].width=25
    width=max((1+2*len(c.cases) for c in combinations),default=1)
    for i,c in enumerate(combinations,1):
        ws.cell(i,1,c.name).data_type='s'
        for j,(cid,coef) in enumerate(c.cases):
            ws.cell(i,2+2*j,int(cid)).number_format='0'
            ws.cell(i,3+2*j,float(coef)).number_format='0.############'
        for col in range(1,width+1):
            cell=ws.cell(i,col)
            cell.font=Font(name='Calibri',size=11,color='17324D')
            cell.alignment=Alignment(horizontal='left' if col==1 else 'right')
            if i%2==0:cell.fill=PatternFill('solid',fgColor='EDF3FA')
        ws.row_dimensions[i].height=20
    for col in range(2,width+1):ws.column_dimensions[get_column_letter(col)].width=13
    ws.freeze_panes='B1'
    wb.defined_names.add(DefinedName('COPIAR_ROBOT',attr_text=f"'ROBOT'!$A$1:${get_column_letter(width)}${max(1,len(combinations))}"))
    info=wb.create_sheet('LEIA-ME')
    info.column_dimensions['A'].width=110
    lines=[
        'COPIAR COMBINAÇÕES PARA O ROBOT',
        '1. Na aba ROBOT, selecione as células preenchidas desde A1 e copie.',
        'Atalho: na Caixa de Nome do Excel (à esquerda da barra de fórmulas), digite COPIAR_ROBOT e pressione Enter.',
        '2. No Robot, cole a partir da coluna Nome, na linha destinada às novas combinações.',
        '3. Confira os números dos casos e os coeficientes; ajuste a classificação ELU/ELS no Robot.',
        'A aba ROBOT não tem cabeçalhos nem a coluna de numeração automática de combinações.',
        'Cada número de caso é exportado como cadastrado, mesmo que a sequência tenha lacunas (ex.: 1, 2, 4).',
        'Casos e coeficientes são números em células separadas. O separador decimal é exibido conforme o Excel.',
        'O aplicativo não verifica a existência desses números no modelo do Robot.',
    ]
    for i,line in enumerate(lines,1):
        cell=info.cell(i,1,line);cell.font=Font(name='Calibri',size=12,color='17324D',bold=i==1)
        cell.alignment=Alignment(wrap_text=True,vertical='center');info.row_dimensions[i].height=40 if i>1 else 32
    out=BytesIO();wb.save(out)
    return out.getvalue()
