import unittest
from copy import deepcopy
from io import BytesIO
from decimal import Decimal
from openpyxl import load_workbook
from engine import example_project, generate, sync_global_families
from project_io import dump_project,load_project,robot_tsv
from excel_export import robot_xlsx
from streamlit.testing.v1 import AppTest
from pathlib import Path

class PY02Tests(unittest.TestCase):
    def test_global_families_overrides_old_action_flags(self):
        p=example_project();p['families']=['ELUN','ELSF']
        for a in p['actions']:a['families']=[]
        self.assertEqual(generate(p).counts,{'ELUN':28,'ELSF':10})
    def test_legacy_schema_import_preserves_numbers_and_links(self):
        p=example_project();p.update(schema=1,app_version='PY01',families=['ELUN'])
        p['actions'][1]['case']=24
        p['actions'][3]['origin']=p['actions'][2]['origin']
        q=load_project(dump_project(p))
        self.assertEqual(q['schema'],2)
        self.assertEqual(q['actions'][1]['case'],24)
        self.assertEqual(q['actions'][3]['origin'],q['actions'][2]['origin'])
        self.assertEqual(q['actions'][0]['families'],['ELUN'])
    def platform(self):
        p=example_project();p['actions']=p['actions'][:3];p['families']=['ELUN','ELSF']
        p['actions'][1].update(type='IND',profile=None,group=None,compatibility=None,case=2)
        p['actions'][2].update(type='OUT',profile='USO1',group=None,compatibility=None,case=4)
        return p
    def test_case4_kept_in_all_exports(self):
        p=self.platform();r=generate(p)
        self.assertEqual(r.counts,{'ELUN':8,'ELSF':2})
        wb=load_workbook(BytesIO(robot_xlsx(r.combinations)))
        ws=wb['ROBOT'];self.assertEqual(ws['F1'].value,4);self.assertEqual(ws['F1'].data_type,'n')
        self.assertEqual(ws['G1'].value,1.5);self.assertEqual(ws['A1'].value,'EX_ELUN_00001')
        self.assertEqual(ws.max_row,10);self.assertEqual(ws.max_column,7)
        self.assertEqual(wb.defined_names['COPIAR_ROBOT'].attr_text,"'ROBOT'!$A$1:$G$10")
        for row,c in zip(ws.iter_rows(values_only=True),r.combinations):
            pairs=tuple((row[i],Decimal(str(row[i+1]))) for i in range(1,len(row),2) if row[i] is not None)
            self.assertEqual(pairs,c.cases)
        for line,c in zip(robot_tsv(r.combinations).splitlines(),r.combinations):
            fields=line.split('\t')
            self.assertEqual([int(x) for x in fields[1::2] if x],[cid for cid,coef in c.cases])
    def test_favorable_without_note_for_all_standards(self):
        for std in ('NBR 8800:2024','NBR 14762:2010','NBR 8681:2025'):
            p=self.platform();p['standard']=std;p['actions'][0].update(g_effect='favorable',notes='')
            self.assertTrue(generate(p).combinations)
    def app(self):
        at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run()
        next(b for b in at.button if b.label=='Carregar exemplo didático').click().run()
        return at
    def click(self,at,label):
        next(x for x in at.button if x.label==label).click().run()
        self.assertFalse(at.exception)
    def test_no_repeated_fields_or_limits(self):
        at=self.app()
        self.assertFalse(any(x.label=='Participa das combinações' for x in at.multiselect))
        self.assertFalse(any('Origem' in x.label for x in at.text_input))
        self.assertFalse(any('Máximo' in x.label for x in at.number_input))
        next(x for x in at.multiselect if x.label=='Combinações a gerar').set_value(['ELSF']).run()
        self.click(at,'Gerar combinações')
        self.assertEqual(at.metric[0].value,'10')
        self.assertTrue(all(a['families']==['ELSF'] for a in at.session_state['project']['actions']))
    def test_new_action_auto_origin_and_families(self):
        at=self.app()
        next(x for x in at.text_input if x.label=='Nome do carregamento').set_value('Permanente adicional')
        next(x for x in at.selectbox if x.label=='Tipo de carregamento').set_value('ACO').run()
        self.click(at,'Adicionar ao cadastro')
        a=at.session_state['project']['actions'][-1]
        self.assertEqual(a['case'],7);self.assertEqual(a['origin'],'CASO_7')
        self.assertEqual(a['families'],at.session_state['project']['families'])
    def test_favorable_edit_without_note(self):
        at=self.app();self.click(at,'Editar ação')
        next(x for x in at.selectbox if x.label=='Ponderações a considerar').set_value('favorable')
        next(x for x in at.text_area if x.label=='Observações').set_value('').run()
        self.click(at,'Salvar ação');self.click(at,'Gerar combinações')
        self.assertFalse(at.error)
        self.assertEqual(at.metric[0].value,'40')
    def test_advanced_link_survives_edit(self):
        at=self.app()
        next(x for x in at.text_input if x.label=='Nome do carregamento').set_value('Parte do PP')
        next(x for x in at.selectbox if x.label=='Tipo de carregamento').set_value('ACO').run()
        next(x for x in at.checkbox if x.label=='Vincular a outro caso da mesma ação física').check().run()
        next(x for x in at.selectbox if x.label=='Caso ao qual vincular').set_value(1).run()
        self.click(at,'Adicionar ao cadastro');self.click(at,'Gerar combinações')
        self.assertFalse(at.error)
        self.assertEqual(at.metric[0].value,'54')
        actions=at.session_state['project']['actions']
        self.assertEqual(actions[-1]['origin'],actions[0]['origin'])

if __name__=='__main__':unittest.main()
