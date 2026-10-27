import unittest
from copy import deepcopy
from io import BytesIO
from decimal import Decimal as D
from zipfile import ZipFile
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from engine import Combination, FAMILIES, InputError, example_project, generate
from project_io import combination_type, robot_tsv
from excel_export import robot_xlsx
from word_export import coefficient_text, combination_expression, combinations_docx

class ExportTests(unittest.TestCase):
    def test_all_family_types(self):
        for family in FAMILIES:
            self.assertEqual(combination_type(family), 'ELU' if family.startswith('ELU') else 'ELS')
            self.assertEqual(combination_type(family,'en'), 'ULS' if family.startswith('ELU') else 'SLS')
        with self.assertRaises(InputError):combination_type('WRONG')
        with self.assertRaises(InputError):combination_type('ELUN','wrong')
    def test_types_do_not_depend_on_names(self):
        c=Combination('ELS_nome_enganoso','ELUN',((4,D('1.25')),))
        self.assertEqual(robot_tsv([c]),'ELS_nome_enganoso\tELU\t4\t1,25')
        self.assertEqual(robot_tsv([c],False,'en'),'ELS_nome_enganoso\tULS\t4\t1.25')
    def test_xlsx_and_tsv_alignment(self):
        rows=generate(example_project()).combinations
        for language in ('pt','en'):
            ws=load_workbook(BytesIO(robot_xlsx(rows,language)))['ROBOT']
            lines=robot_tsv(rows,type_language=language).splitlines()
            for i,(c,line) in enumerate(zip(rows,lines),1):
                fields=line.split('\t')
                self.assertEqual(fields[:2],[c.name,combination_type(c.family,language)])
                self.assertEqual(ws.cell(i,2).value,fields[1])
                for j,(cid,coef) in enumerate(c.cases):
                    self.assertEqual(int(fields[2+2*j]),cid)
                    self.assertEqual(D(fields[3+2*j].replace(',','.')),coef)
                    self.assertEqual(ws.cell(i,3+2*j).value,cid)
                    self.assertEqual(D(str(ws.cell(i,4+2*j).value)),coef)
    def test_word_exact_precision_and_signs(self):
        c=Combination('X','ELUN',((1,D('1.25')),(9,D('-0.375')),(8,D('0'))),detail=({'case':1,'name':'PP'},{'case':9,'name':'Vento -X'},{'case':8,'name':'SC'}))
        self.assertEqual(combination_expression(c),'1,25PP - 0,375Vento -X')
        self.assertEqual(coefficient_text(D('1')),'1,00')
        self.assertEqual(coefficient_text(D('0.000001')),'0,000001')
        with self.assertRaises(InputError):coefficient_text(D('NaN'))
    def test_word_missing_name_rejected(self):
        with self.assertRaises(InputError):combination_expression(Combination('X','ELUN',((7,D('1')),)))
    def test_word_selected_rows_order_and_no_equations(self):
        rows=[c for c in generate(example_project()).combinations if c.family=='ELSF']
        before=deepcopy(rows)
        data=combinations_docx(rows,'Projeto teste')
        doc=Document(BytesIO(data));strings=[p.text for p in doc.paragraphs]
        expressions=[t for t in strings if t[:1].isdigit()]
        self.assertEqual(len(expressions),len(rows))
        for i,(line,c) in enumerate(zip(expressions,rows),1):
            self.assertEqual(line,f'{i}- {combination_expression(c)}')
        self.assertEqual(rows,before)
        self.assertNotIn(FAMILIES['ELUN'],strings)
        with ZipFile(BytesIO(data)) as z:
            self.assertFalse(any(n.startswith('word/media/') for n in z.namelist()))
            self.assertNotIn(b'oMath',z.read('word/document.xml'))
    def test_word_xml_controls_do_not_break_exports(self):
        c=Combination('X','ELUN',((1,D('1.25')),),detail=({'case':1,'name':'PP\x00teste'},))
        doc=Document(BytesIO(combinations_docx([c],'Projeto\x00teste')))
        self.assertIn('1- 1,25PPteste',[p.text for p in doc.paragraphs])
        self.assertIn('Projetoteste',[p.text for p in doc.paragraphs])
    def test_empty_exports(self):
        self.assertEqual(robot_tsv([]),'')
        self.assertTrue(combinations_docx([]))
    def test_app_export_language_and_stale_results(self):
        at=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=30).run()
        next(b for b in at.button if b.label=='Carregar exemplo didático').click().run()
        next(b for b in at.button if b.label=='Gerar combinações').click().run()
        self.assertFalse(at.exception)
        self.assertIn('\tELU\t',at.code[0].value)
        next(r for r in at.radio if r.label=='Tipo de combinação no Robot').set_value('en').run()
        self.assertIn('\tULS\t',at.code[0].value)
        labels=[b.label for b in at.get('download_button')]
        self.assertIn('Baixar Word para MC (.docx)',labels)
        next(x for x in at.multiselect if x.label=='Famílias a copiar / baixar').set_value(['ELSF']).run()
        self.assertEqual(len(at.code[0].value.splitlines()),10)
        self.assertNotIn('\tULS\t',at.code[0].value)
        next(x for x in at.text_input if x.label=='Projeto').set_value('Projeto alterado').run()
        self.assertFalse(at.code)
        self.assertNotIn('Baixar Word para MC (.docx)',[b.label for b in at.get('download_button')])

if __name__=='__main__':unittest.main()
